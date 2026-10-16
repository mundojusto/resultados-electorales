"""Tests del volcado de avales de la hoja a la web (herramientas/actualizar_avales.py)."""
from __future__ import annotations

import actualizar_avales as aa
import openpyxl
import pytest

CABECERA = ["Censo CER+CERA", "Avales legales (0,1%)", "Objetivo +20%", "Objetivo +30%",
            "Planificadas", "Conseguidas", "% sobre legal"]


def libro(pestanas: dict[str, object], cabecera_f="Conseguidas"):
    wb = openpyxl.Workbook()
    wb.active.title = "Resumen"
    wb.active["F3"] = "Planificados"
    for nombre, f4 in pestanas.items():
        ws = wb.create_sheet(nombre)
        for i, h in enumerate(CABECERA, start=1):
            ws.cell(row=3, column=i, value=h)
        ws["F3"] = cabecera_f
        ws["F4"] = f4
    return wb


def provincia(nombre, censo=100_000, papel=None, estado="No iniciado"):
    return {"provincia": nombre, "censo": censo, "avales_papel": papel,
            "avales_boreal": None, "estado": estado}


@pytest.mark.parametrize("valor,esperado", [
    (1885, 1885), (12.0, 12), ("1.379", 1379), ("-", 0), ("", 0), (None, 0), ("n/d", None),
])
def test_a_numero(valor, esperado):
    assert aa.a_numero(valor) == esperado


def test_lee_pestanas_y_alias():
    cifras, avisos = aa.leer_conseguidas(libro({"Madrid": 1885, "Araba-Álava": "62"}))
    assert cifras == {"MADRID": 1885, "ARABA/ALAVA": 62}
    assert avisos == []


def test_ignora_pestanas_con_otra_cabecera():
    cifras, _ = aa.leer_conseguidas(libro({"Soria": 98}, cabecera_f="Otra cosa"))
    assert cifras == {}


def test_avisa_si_f4_no_es_numero():
    cifras, avisos = aa.leer_conseguidas(libro({"Soria": "#REF!"}))
    assert cifras == {}
    assert "Soria" in avisos[0]


@pytest.mark.parametrize("recogidos,estado", [
    (0, "No iniciado"), (10, "En recogida"), (80, "Cerca del mínimo"),
    (100, "Mínimo alcanzado"), (120, "Margen alcanzado"),
])
def test_estado_segun_minimo(recogidos, estado):
    # Censo 100.000 → mínimo legal 100, margen +20 % = 120.
    assert aa.estado_para(recogidos, 100_000, "En recogida") == estado


def test_no_pisa_estados_manuales():
    assert aa.estado_para(500, 100_000, "Presentado") == "Presentado"


def test_aplicar_solo_toca_provincias_de_la_hoja():
    ps = [provincia("Madrid"), provincia("Lugo", papel=3, estado="En recogida")]
    cambios = aa.aplicar(ps, {"MADRID": 150})
    assert ps[0]["avales_papel"] == 150 and ps[0]["estado"] == "Margen alcanzado"
    assert ps[1]["avales_papel"] == 3  # sin pestaña: se queda como estaba
    assert cambios == ["Madrid: — → 150"]


def test_aplicar_sin_cambios():
    ps = [provincia("Madrid", papel=150, estado="Margen alcanzado")]
    assert aa.aplicar(ps, {"MADRID": 150}) == []


def test_lectura_de_manana_cuenta_como_cierre_del_dia_anterior():
    from datetime import date, datetime
    assert aa.dia_de_cierre(datetime(2026, 10, 17, 8, 23)) == date(2026, 10, 16)
    assert aa.dia_de_cierre(datetime(2026, 10, 17, 18, 0)) == date(2026, 10, 17)


def test_registrar_dia_sustituye_el_mismo_dia_y_ordena():
    from datetime import date
    ps = [{"codigo": "28", "avales_papel": 10, "avales_boreal": 2},
          {"codigo": "27", "avales_papel": None, "avales_boreal": None}]
    h = [{"fecha": "2026-10-16", "total": 1, "prov": {"28": 1}}]
    assert aa.registrar_dia(h, date(2026, 10, 15), ps)
    assert aa.registrar_dia(h, date(2026, 10, 16), ps)
    assert h == [{"fecha": "2026-10-15", "total": 12, "prov": {"28": 12}},
                 {"fecha": "2026-10-16", "total": 12, "prov": {"28": 12}}]


def test_registrar_dia_fuera_de_la_recogida():
    from datetime import date
    h: list[dict] = []
    assert not aa.registrar_dia(h, date(2026, 10, 27), [])
    assert h == []
