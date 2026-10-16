"""Vuelca los avales recogidos de la hoja de Google a la web (pestaña Avales 29-N).

Fuente: la hoja «MJ_Avales_Generales_2026». Cada pestaña provincial tiene en la
fila 3 la cabecera y en la fila 4 las cifras; F4 («Conseguidas») es el total de
avales recogidos en la provincia (sede y brigada). La hoja es la misma que
alimenta el panel «Recogida de avales 29-N», así que web y panel cuadran.

Se descarga el libro entero como XLSX con el enlace público (no hace falta
ninguna clave) y se actualizan en `web/src/avales/provincias-29n.json`:

- `avales_papel`: el valor de F4 (la hoja no separa papel y Boreal, así que el
  total va aquí y `avales_boreal` queda como esté);
- `estado`: se deduce de la cifra, salvo que la Comisión ya haya marcado
  «Revisado» o «Presentado», que nunca se pisan.

Además guarda en `web/src/avales/historial-avales.json` lo acumulado en cada
provincia al cierre de cada día, para la gráfica de evolución. Una lectura
hecha antes de las 12:00 cuenta como cierre del día anterior: por la mañana
apenas se ha recogido nada todavía, y así la ejecución diaria de las 8:23 deja
cada punto con su fecha real. Una lectura posterior del mismo día se pisa con
la siguiente.

Si una pestaña falta o su cabecera ya no dice lo esperado, esa provincia no se
toca. Si no se puede leer ninguna, el script falla sin escribir nada.

Uso:
    python herramientas/actualizar_avales.py            # descarga y escribe
    python herramientas/actualizar_avales.py --xlsx libro.xlsx   # fichero local
"""
from __future__ import annotations

import argparse
import io
import json
import math
import sys
import unicodedata
import urllib.request
from datetime import date, datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl

HOJA_ID = "1hTGTJzaw7UaevRDlcjivWc-rdgPYkEGBlsFZ11wJflg"
URL_XLSX = f"https://docs.google.com/spreadsheets/d/{HOJA_ID}/export?format=xlsx"

RAIZ = Path(__file__).resolve().parent.parent
JSON_PROVINCIAS = RAIZ / "web/src/avales/provincias-29n.json"
JSON_ACTUALIZACION = RAIZ / "web/src/avales/avales-actualizacion.json"
JSON_HISTORIAL = RAIZ / "web/src/avales/historial-avales.json"

# Días que entran en el histórico: los de la recogida.
PRIMER_DIA = date(2026, 10, 6)
ULTIMO_DIA = date(2026, 10, 26)

# Pestañas cuyo nombre no coincide con el de la web.
ALIAS = {"ARABA-ALAVA": "ARABA/ALAVA"}

# Estados que pone a mano la Comisión Electoral: el script no los cambia.
ESTADOS_MANUALES = {"Revisado", "Presentado"}


def normaliza(texto: str) -> str:
    sin_tildes = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    return " ".join(sin_tildes.upper().split())


def a_numero(valor) -> int | None:
    """Convierte una celda en entero. «-» o vacío es 0; texto no numérico, None."""
    if valor is None:
        return 0
    if isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return max(0, round(valor))
    texto = str(valor).strip()
    if texto in ("", "-", "—"):
        return 0
    texto = texto.replace(".", "").replace(" ", "").replace(" ", "")
    return int(texto) if texto.isdigit() else None


def minimo_legal(censo: int) -> int:
    return math.ceil(censo / 1000)


def objetivo(censo: int, margen: float) -> int:
    return math.ceil(minimo_legal(censo) * (1 + margen) - 1e-9)


def estado_para(recogidos: int, censo: int, actual: str) -> str:
    if actual in ESTADOS_MANUALES:
        return actual
    minimo = minimo_legal(censo)
    if recogidos >= objetivo(censo, 0.2):
        return "Margen alcanzado"
    if recogidos >= minimo:
        return "Mínimo alcanzado"
    if recogidos >= 0.8 * minimo:
        return "Cerca del mínimo"
    if recogidos > 0:
        return "En recogida"
    return "No iniciado"


def leer_libro(origen: Path | None):
    if origen is not None:
        return openpyxl.load_workbook(origen, data_only=True, read_only=True)
    with urllib.request.urlopen(URL_XLSX, timeout=120) as r:
        datos = r.read()
    return openpyxl.load_workbook(io.BytesIO(datos), data_only=True, read_only=True)


def leer_conseguidas(libro) -> tuple[dict[str, int], list[str]]:
    """Devuelve {nombre normalizado de pestaña: F4} y la lista de avisos."""
    cifras: dict[str, int] = {}
    avisos: list[str] = []
    for hoja in libro.worksheets:
        filas = list(hoja.iter_rows(min_row=3, max_row=4, max_col=6, values_only=True))
        if len(filas) < 2 or len(filas[0]) < 6:
            continue
        cabecera, valores = filas[0], filas[1]
        if normaliza(str(cabecera[5] or "")) != "CONSEGUIDAS":
            continue  # Resumen, BOREAL, HojasFirmas u otra pestaña sin ese formato
        n = a_numero(valores[5])
        if n is None:
            avisos.append(f"{hoja.title}: F4 no es un número ({valores[5]!r})")
            continue
        nombre = normaliza(hoja.title)
        cifras[ALIAS.get(nombre, nombre)] = n
    return cifras, avisos


def aplicar(provincias: list[dict], cifras: dict[str, int]) -> list[str]:
    """Escribe las cifras en `provincias` y devuelve las líneas de cambios."""
    cambios = []
    for p in provincias:
        n = cifras.get(normaliza(p["provincia"]))
        if n is None:
            continue
        estado = estado_para(n, p["censo"], p["estado"])
        if p["avales_papel"] != n or p["estado"] != estado:
            antes = p["avales_papel"]
            cambios.append(f"{p['provincia']}: {antes if antes is not None else '—'} → {n}")
            p["avales_papel"] = n
            p["estado"] = estado
    return cambios


def dia_de_cierre(ahora: datetime) -> date:
    """Día cuyo cierre refleja una lectura hecha en `ahora` (hora de Madrid)."""
    return ahora.date() - timedelta(days=1) if ahora.hour < 12 else ahora.date()


def registrar_dia(historial: list[dict], dia: date, provincias: list[dict]) -> bool:
    """Guarda (o sustituye) el punto de `dia`. False si cae fuera de la recogida."""
    if not PRIMER_DIA <= dia <= ULTIMO_DIA:
        return False
    prov = {}
    for p in provincias:
        n = (p["avales_papel"] or 0) + (p["avales_boreal"] or 0)
        if n:
            prov[p["codigo"]] = n
    punto = {"fecha": dia.isoformat(), "total": sum(prov.values()), "prov": prov}
    historial[:] = [h for h in historial if h["fecha"] != punto["fecha"]] + [punto]
    historial.sort(key=lambda h: h["fecha"])
    return True


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--xlsx", type=Path, help="Leer de un XLSX local en vez de descargarlo")
    args = ap.parse_args(argv)

    cifras, avisos = leer_conseguidas(leer_libro(args.xlsx))
    for a in avisos:
        print(f"::warning::{a}")
    if not cifras:
        print("::error::No se ha podido leer ninguna provincia de la hoja.")
        return 1

    provincias = json.loads(JSON_PROVINCIAS.read_text(encoding="utf-8"))
    conocidas = {normaliza(p["provincia"]) for p in provincias}
    for nombre in sorted(set(cifras) - conocidas):
        print(f"::warning::La pestaña «{nombre}» no corresponde a ninguna provincia de la web.")

    cambios = aplicar(provincias, cifras)
    total = sum(n for nombre, n in cifras.items() if nombre in conocidas)
    print(f"{len(cifras)} provincias leídas · {total} avales en total · {len(cambios)} cambios")
    for c in cambios:
        print(f"  {c}")
    if cambios:
        JSON_PROVINCIAS.write_text(
            json.dumps(provincias, ensure_ascii=False, indent=1), encoding="utf-8"
        )
    # La fecha se renueva aunque no cambie nada: la web dice cuándo se comprobó.
    ahora = datetime.now(ZoneInfo("Europe/Madrid"))
    JSON_ACTUALIZACION.write_text(
        json.dumps({"actualizado": ahora.isoformat(timespec="minutes")}, ensure_ascii=False,
                   indent=1) + "\n",
        encoding="utf-8",
    )

    historial = json.loads(JSON_HISTORIAL.read_text(encoding="utf-8"))
    dia = dia_de_cierre(ahora)
    if registrar_dia(historial, dia, provincias):
        # Un día por línea: los diffs diarios se leen de un vistazo.
        lineas = ",\n".join(json.dumps(h, ensure_ascii=False) for h in historial)
        JSON_HISTORIAL.write_text(f"[\n{lineas}\n]\n", encoding="utf-8")
        print(f"Histórico: cierre del {dia.isoformat()} guardado.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
